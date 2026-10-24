"""reports.service with the db layer faked: who may download, which sheets,
and paging through every row. The SQL is in tests/db/, the whole file
against Kotlin's in tests/contract/ (K-42)."""

import io
from datetime import date
from decimal import Decimal
from typing import Any
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from src import db
from src.actor import Actor
from src.db.reports import FarmExportRow
from src.exceptions import Forbidden, Invalid
from src.reports import service

SESSION: Any = object()


def _actor(perms: frozenset[str]) -> Actor:
    return Actor(user_id=uuid4(), roles=frozenset(), perms=perms, channel="api")


RESEARCHER = _actor(frozenset({"read:report:all"}))


def _farm(name: str) -> FarmExportRow:
    return FarmExportRow(
        farm_name=name,
        found_date=date(2020, 1, 1),
        total_area=Decimal("1"),
        province_name_en=None,
        district_name_en=None,
        subdistrict_name_en=None,
        contact_name=None,
        phone_number=None,
        line=None,
        facebook=None,
    )


@pytest.fixture
def fake_db(monkeypatch: pytest.MonkeyPatch) -> dict[str, AsyncMock]:
    fakes = {
        "farm_export_page": AsyncMock(return_value=[]),
        "harvest_export_page": AsyncMock(return_value=[]),
    }
    for name, mock in fakes.items():
        monkeypatch.setattr(db.reports, name, mock)
    return fakes


def _sheets(content: bytes) -> dict[str, list[tuple[Any, ...]]]:
    wb = load_workbook(io.BytesIO(content), read_only=True)
    return {ws.title: list(ws.iter_rows(values_only=True)) for ws in wb.worksheets}


async def test_no_sheets_means_both_in_kotlins_order(fake_db: dict[str, AsyncMock]) -> None:
    export = await service.raw_data_xlsx(SESSION, RESEARCHER)

    assert list(_sheets(export.content)) == ["Farm Registry", "Raw Harvest Data"]
    assert export.filename.startswith("raw_data_") and export.filename.endswith(".xlsx")


@pytest.mark.parametrize(
    ("asked", "sheets"),
    [
        (["HARVEST"], ["Raw Harvest Data"]),
        (["farm"], ["Farm Registry"]),
        (["HARVEST,FARM"], ["Raw Harvest Data", "Farm Registry"]),
        (["FARM", "HARVEST", "FARM"], ["Farm Registry", "Raw Harvest Data"]),
    ],
)
async def test_sheets_asked_for_in_that_order(
    fake_db: dict[str, AsyncMock], asked: list[str], sheets: list[str]
) -> None:
    export = await service.raw_data_xlsx(SESSION, RESEARCHER, asked)
    assert list(_sheets(export.content)) == sheets


@pytest.mark.parametrize("asked", [["NOPE"], ["FARM,NOPE"], [" , "]])
async def test_unknown_sheet_is_invalid(fake_db: dict[str, AsyncMock], asked: list[str]) -> None:
    with pytest.raises(Invalid) as e:
        await service.raw_data_xlsx(SESSION, RESEARCHER, asked)
    assert "sheets" in e.value.fields
    fake_db["farm_export_page"].assert_not_awaited()


async def test_export_is_for_researchers(fake_db: dict[str, AsyncMock]) -> None:
    with pytest.raises(Forbidden):
        await service.raw_data_xlsx(SESSION, _actor(frozenset({"read:farm:own"})))
    fake_db["farm_export_page"].assert_not_awaited()


async def test_pages_until_an_empty_page(
    fake_db: dict[str, AsyncMock], monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(service, "PAGE_SIZE", 2)
    fake_db["farm_export_page"].side_effect = [[_farm("a"), _farm("b")], [_farm("c")], []]

    export = await service.raw_data_xlsx(SESSION, RESEARCHER, ["FARM"])

    rows = _sheets(export.content)["Farm Registry"]
    assert [r[0] for r in rows[1:]] == ["a", "b", "c"]
    offsets = [c.kwargs["offset"] for c in fake_db["farm_export_page"].await_args_list]
    assert offsets == [0, 2, 4]
