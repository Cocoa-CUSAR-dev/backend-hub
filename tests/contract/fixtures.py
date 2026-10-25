"""Shared by the recorder (record_web.py) and the replay test (test_web_contract.py).

A fixture file is one Kotlin route (one K-row of api-migration-map.xlsx) with
the requests recorded against it. Replaying a case sends the same request to
backend-hub and checks the answer means the same thing, after the differences
the migration makes on purpose:

- Kotlin wraps every body in {"value": ..., "error": ...}; the hub doesn't.
- Kotlin answers bad input with 400; the hub with 422 (CONVENTIONS.md §6).
- Error bodies are never compared, only the status: the hub's {"detail"} is
  a different text from Kotlin's `error`.

Values that change on every run (ids of rows a POST created, "now", tokens)
are masked on both sides before comparing.
"""

from __future__ import annotations

import json
import re
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

HERE = Path(__file__).parent
WEB = HERE / "web"
SEED_DIR = WEB / "seed"
# Load order; a fixture names the files it needs and they load in this order.
# "reference" (rows the migrations insert) is always loaded, first.
SEED_ORDER = ("reference", "base", "places", "forms", "answers")

# Kotlin status -> the status backend-hub answers for the same case.
STATUS_MAP = {400: 422}

# Who sends a request. Ids match tests/contract/web/seed/base.sql.
ACTORS = {
    "admin": ("20000000-0000-4000-8000-000000000001", "contract_admin"),
    "researcher": ("20000000-0000-4000-8000-000000000002", "contract_researcher"),
    "farmer_1": ("20000000-0000-4000-8000-000000000003", "contract_farmer_1"),
    "farmer_2": ("20000000-0000-4000-8000-000000000004", "contract_farmer_2"),
    "processor": ("20000000-0000-4000-8000-000000000005", "contract_processor"),
    "no_role": ("20000000-0000-4000-8000-000000000006", "contract_no_role"),
    "collector": ("20000000-0000-4000-8000-000000000007", "contract_collector"),
}

_UUID = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", re.I)
_UUID_ANY = re.compile(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}", re.I)
_DATETIME = re.compile(r"^(\d{4}-\d{2}-\d{2})[T ]\d{2}:\d{2}")
_JWT = re.compile(r"^eyJ[\w-]+\.[\w-]+\.[\w-]+$")


@dataclass
class Case:
    name: str
    status: int
    method: str
    path: str
    actor: str | None = None  # a key of ACTORS; None = no token
    # Kotlin path parameters by name ({"formId": "..."}), so the replay can
    # fill the hub's template (its "params" maps hub name -> Kotlin name).
    params: dict[str, str] = field(default_factory=dict)
    query: dict[str, str] = field(default_factory=dict)
    body: Any = None
    response: Any = None
    # What the hub answers where it differs on purpose; recorded with the case
    # where Kotlin is wrong, and added by a porting PR where the hub changes:
    #   {"status": 200, "ignore": ["*.formId"], "compare": "status", "unordered": true}
    hub: dict[str, Any] = field(default_factory=dict)
    service_key: bool = False


@dataclass
class Fixture:
    kid: str
    after: str
    kotlin: dict[str, str]
    # {"method", "path", "query"?, "params"?} or None when the route becomes an
    # in-process call (K-50, K-51); then only service tests use the cases.
    hub: dict[str, Any] | None
    seeds: list[str]
    recorded_on: str
    cases: list[Case]
    note: str = ""

    @property
    def path(self) -> Path:
        return WEB / f"{self.kid}.json"


def load(path: Path) -> Fixture:
    raw = json.loads(path.read_text(encoding="utf-8"))
    cases = [Case(**c) for c in raw.pop("cases")]
    return Fixture(cases=cases, **raw)


def all_fixtures() -> list[Fixture]:
    return [load(p) for p in sorted(WEB.glob("K-*.json"))]


def dump(fx: Fixture) -> str:
    data = {
        "kid": fx.kid,
        "after": fx.after,
        "note": fx.note,
        "kotlin": fx.kotlin,
        "hub": fx.hub,
        "seeds": fx.seeds,
        "recorded_on": fx.recorded_on,
        "cases": [
            {
                k: v
                for k, v in c.__dict__.items()
                if v not in (None, {}, False) or k in ("status", "body")
            }
            for c in fx.cases
        ],
    }
    return json.dumps(data, ensure_ascii=False, indent=2) + "\n"


# ---- seeds ------------------------------------------------------------------


def seed_sql(names: Iterable[str]) -> list[str]:
    wanted = {"reference", *names}
    return [(SEED_DIR / f"{n}.sql").read_text(encoding="utf-8") for n in SEED_ORDER if n in wanted]


def seed_tables(names: Iterable[str]) -> list[str]:
    """Every table the named seed files insert into, in load order."""
    found: list[str] = []
    for sql in seed_sql(names):
        for t in re.findall(r"INSERT INTO ([a-z_]+\.[a-z_]+)", sql):
            if t not in found:
                found.append(t)
    return found


def seed_ids() -> set[str]:
    ids: set[str] = set()
    for n in SEED_ORDER:
        ids |= {
            m.lower()
            for m in _UUID_ANY.findall((SEED_DIR / f"{n}.sql").read_text(encoding="utf-8"))
        }
    return ids


# ---- masking and comparing --------------------------------------------------


def known_ids(case: Case) -> set[str]:
    """Ids a response may show as they are: seeded ones and ones the request named."""
    request = json.dumps([case.path, case.query, case.body], ensure_ascii=False)
    return seed_ids() | {m.lower() for m in _UUID_ANY.findall(request)}


def mask(value: Any, known: set[str], today: date) -> Any:
    if isinstance(value, dict):
        return {k: mask(v, known, today) for k, v in value.items()}
    if isinstance(value, list):
        return [mask(v, known, today) for v in value]
    if isinstance(value, str):
        if _UUID.match(value) and value.lower() not in known:
            return "<new-id>"
        if _JWT.match(value):
            return "<token>"
        m = _DATETIME.match(value)
        # A time from the last day is "now" (a row the request created).
        if m and date.fromisoformat(m.group(1)) >= today - timedelta(days=1):
            return "<now>"
    return value


def _drop(value: Any, parts: list[str]) -> Any:
    if not parts:
        return None
    head, rest = parts[0], parts[1:]
    if isinstance(value, list) and head == "*":
        return [_drop(v, rest) for v in value] if rest else value
    if isinstance(value, dict) and head in value:
        if rest:
            return {**value, head: _drop(value[head], rest)}
        return {k: v for k, v in value.items() if k != head}
    return value


def without(value: Any, paths: Iterable[str]) -> Any:
    """Remove dotted paths (`items.*.formId`) the hub adds or changes on purpose."""
    for p in paths:
        value = _drop(value, p.split("."))
    return value


def expected_status(case: Case) -> int:
    return int(case.hub.get("status", STATUS_MAP.get(case.status, case.status)))


@dataclass
class Diff:
    ok: bool
    message: str = ""


def compare(case: Case, status: int, body: Any, recorded_on: date, today: date) -> Diff:
    """Kotlin's recorded answer (masked as of `recorded_on`) vs the hub's (as of `today`)."""
    want = expected_status(case)
    if status != want:
        return Diff(False, f"status {status}, expected {want} (Kotlin answered {case.status})")
    # Error answers are compared by status only, on either side.
    if case.status >= 400 or want >= 400 or case.hub.get("compare") == "status":
        return Diff(True)
    known = known_ids(case)
    ignore = case.hub.get("ignore", [])
    expected = without(mask(case.response, known, recorded_on), ignore)
    actual = without(mask(body, known, today), ignore)
    if case.hub.get("unordered"):
        expected, actual = _sorted(expected), _sorted(actual)
    if expected == actual:
        return Diff(True)
    return Diff(False, "body differs:\n" + "\n".join(_diff_lines(expected, actual)))


def _sorted(value: Any) -> Any:
    """Lists in a stable order, for answers whose order Kotlin never fixed."""
    if isinstance(value, dict):
        return {k: _sorted(v) for k, v in value.items()}
    if isinstance(value, list):
        return sorted(
            (_sorted(v) for v in value),
            key=lambda v: json.dumps(v, sort_keys=True, ensure_ascii=False),
        )
    return value


def _diff_lines(a: Any, b: Any, at: str = "$") -> Iterator[str]:
    if isinstance(a, dict) and isinstance(b, dict):
        for k in sorted(set(a) | set(b), key=str):
            if k not in b:
                yield f"  {at}.{k}: missing in hub (Kotlin: {a[k]!r})"
            elif k not in a:
                yield f"  {at}.{k}: only in hub ({b[k]!r})"
            else:
                yield from _diff_lines(a[k], b[k], f"{at}.{k}")
    elif isinstance(a, list) and isinstance(b, list):
        if len(a) != len(b):
            yield f"  {at}: {len(a)} items in Kotlin, {len(b)} in hub"
        for i, (x, y) in enumerate(zip(a, b, strict=False)):
            yield from _diff_lines(x, y, f"{at}[{i}]")
    elif a != b:
        yield f"  {at}: Kotlin {a!r}, hub {b!r}"


# ---- Kotlin bodies ----------------------------------------------------------


def kotlin_value(body: Any) -> Any:
    """Kotlin's ApiResponse: {"value": ..., "error": ...} -> the value, or the error text."""
    if isinstance(body, dict) and set(body) <= {"value", "error"}:
        return body.get("value") if body.get("error") is None else {"error": body["error"]}
    return body


def xlsx_summary(content: bytes) -> dict[str, list[list[Any]]]:
    """An .xlsx as {sheet: rows}, cells as text, so two exports can be compared."""
    import io

    from openpyxl import load_workbook  # only the raw-data export needs it

    wb = load_workbook(io.BytesIO(content), read_only=True)

    def cell(value: Any) -> str:
        if value is None:
            return ""
        # A number is compared by value: Apache POI writes 4.0 as "4.0" and
        # openpyxl as "4", which Excel shows the same.
        if isinstance(value, int | float) and not isinstance(value, bool):
            return str(float(value))
        return str(value)

    out: dict[str, list[list[Any]]] = {}
    for ws in wb.worksheets:
        out[ws.title] = [[cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
    return out


def to_date(stamp: str) -> date:
    return datetime.fromisoformat(stamp).date()
